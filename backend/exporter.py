import os
import uuid
import shutil
import zipfile
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


def save_validation_results(
    original_file: str,
    python_errors,
    ai_report: str,
    output_dir: str = "temp_uploads",
):
    """
    요구사항 반영 버전

    1) .xlsx 입력
       - 원본을 output_dir로 복사(서식 보존)
       - 오류 셀 빨간색 하이라이트
       - '오류리포트', 'AI_REPORT' 시트 생성
       - 결과는 .xlsx 1개로 반환

    2) .xls 입력
       - .xls 원본은 변형/변환하지 않음(형식 유지)
       - 오류리포트/AI_REPORT만 들어있는 report.xlsx 생성
       - (원본 .xls + report.xlsx) 를 zip으로 묶어 반환
         => "원본 형식 유지 + 오류리포트 제공"을 가장 안정적으로 달성

    반환값: output_path (다운로드할 파일 경로)
      - xlsx 입력이면 *.xlsx
      - xls 입력이면 *.zip
    """
    os.makedirs(output_dir, exist_ok=True)

    ext = os.path.splitext(original_file)[1].lower()
    uid = uuid.uuid4().hex[:10]

    # -----------------------------
    # 내부 유틸
    # -----------------------------
    def _norm(s) -> str:
        return str(s or "").strip().replace("\n", " ")

    def _find_header_row_from_ws(ws, max_scan=50) -> int:
        """워크시트에서 '사원번호'가 포함된 행을 헤더로 간주(1-based)"""
        scan = min(max_scan, ws.max_row)
        for r in range(1, scan + 1):
            row_values = [_norm(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)]
            if any("사원번호" in v for v in row_values):
                return r
        return 1

    def _build_col_map_from_ws(ws, header_row: int) -> dict:
        """컬럼명(문자열) -> 열번호(1-based)"""
        col_map = {}
        for c in range(1, ws.max_column + 1):
            key = ws.cell(row=header_row, column=c).value
            if key is not None:
                col_map[_norm(key)] = c
        return col_map

    def _resolve_col_idx(col_map: dict, col_name: str | None) -> int | None:
        """에러의 column(문자열)을 실제 엑셀 열번호로 매핑(정확 일치 -> 포함 매칭)"""
        if not col_name:
            return None
        col_name = _norm(col_name)

        # 1) 정확 일치
        if col_name in col_map:
            return col_map[col_name]

        # 2) 포함 매칭
        for k, idx in col_map.items():
            if col_name and col_name in k:
                return idx

        return None

    def _find_col_idx_by_keywords(col_map: dict, keywords: list[str]) -> int | None:
        """대표 필드(사원번호/입사일자/기준급여 등) 뽑을 때 키워드로 찾기"""
        # 정확 일치
        for kw in keywords:
            kw = _norm(kw)
            if kw in col_map:
                return col_map[kw]
        # 포함 매칭
        for kw in keywords:
            kw = _norm(kw)
            for k, idx in col_map.items():
                if kw in k:
                    return idx
        return None

    def _safe_sheet_select(wb: openpyxl.Workbook):
        if "(2-2) 재직자 명부" in wb.sheetnames:
            return wb["(2-2) 재직자 명부"]
        if "(2-1) 명부" in wb.sheetnames:
            return wb["(2-1) 명부"]
        return wb.active

    def _delete_sheet_if_exists(wb, name: str):
        if name in wb.sheetnames:
            del wb[name]

    def _autosize(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value is None:
                    continue
                max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 2, 60)

    def _err_msg(err: dict) -> str:
        # 너 rules.py는 msg 사용, 예전 코드 일부는 message를 쓰기도 함
        return err.get("msg") or err.get("message") or ""

    # -----------------------------
    # 케이스 1) .xlsx : 원본에 하이라이트 + 시트 추가
    # -----------------------------
    if ext == ".xlsx":
        result_name = f"validation_{uid}.xlsx"
        output_path = os.path.join(output_dir, result_name)

        try:
            # 1) 원본 복사(서식 보존)
            shutil.copy2(original_file, output_path)
            wb = openpyxl.load_workbook(output_path)

            # 2) 대상 시트/헤더 찾기
            ws = _safe_sheet_select(wb)
            header_row = _find_header_row_from_ws(ws)
            col_map = _build_col_map_from_ws(ws, header_row)

            # 3) 오류 하이라이트
            error_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            for err in (python_errors or []):
                row_idx = err.get("row")
                col_name = err.get("column")
                if not row_idx or not col_name:
                    continue
                try:
                    excel_row = int(row_idx)
                except:
                    continue
                col_idx = _resolve_col_idx(col_map, col_name)
                if col_idx:
                    ws.cell(row=excel_row, column=col_idx).fill = error_fill

            # 4) 오류리포트 시트 생성(원본 행의 주요 값 그대로 복사)
            _delete_sheet_if_exists(wb, "오류리포트")
            report_ws = wb.create_sheet("오류리포트", 0)

            # 대표 필드 열 인덱스 찾기
            c_emp = _find_col_idx_by_keywords(col_map, ["사원번호", "사번", "ID"])
            c_hire = _find_col_idx_by_keywords(col_map, ["입사일자", "입사일", "채용일"])
            c_pay = _find_col_idx_by_keywords(col_map, ["기준급여", "기본급"])
            c_birth = _find_col_idx_by_keywords(col_map, ["생년월일", "출생", "생일"])

            headers = [
                "No", "원본시트", "원본행",
                "사원번호", "생년월일", "입사일자", "기준급여",
                "오류컬럼", "오류메시지"
            ]
            for j, h in enumerate(headers, start=1):
                cell = report_ws.cell(row=1, column=j, value=h)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")

            out_r = 2
            for i, err in enumerate(python_errors or [], start=1):
                row_idx = err.get("row")
                col_name = err.get("column")
                if not row_idx:
                    continue
                try:
                    excel_row = int(row_idx)
                except:
                    continue

                report_ws.cell(out_r, 1, i)
                report_ws.cell(out_r, 2, ws.title)
                report_ws.cell(out_r, 3, excel_row)

                report_ws.cell(out_r, 4, ws.cell(excel_row, c_emp).value if c_emp else None)
                report_ws.cell(out_r, 5, ws.cell(excel_row, c_birth).value if c_birth else None)
                report_ws.cell(out_r, 6, ws.cell(excel_row, c_hire).value if c_hire else None)
                report_ws.cell(out_r, 7, ws.cell(excel_row, c_pay).value if c_pay else None)

                report_ws.cell(out_r, 8, _norm(col_name))
                report_ws.cell(out_r, 9, _err_msg(err))

                out_r += 1

            _autosize(report_ws)

            # 5) AI_REPORT 시트
            _delete_sheet_if_exists(wb, "AI_REPORT")
            ai_ws = wb.create_sheet("AI_REPORT")
            ai_ws["A1"] = "AI 오류 분석 리포트"
            ai_ws["A1"].font = Font(bold=True, size=14)
            ai_ws["A3"] = ai_report or "(리포트 없음)"
            ai_ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
            ai_ws.column_dimensions["A"].width = 120
            ai_ws.row_dimensions[3].height = 600

            wb.save(output_path)
            return output_path

        except Exception as e:
            print("❌ exporter(xlsx) error:", e)
            return None

    # -----------------------------
    # 케이스 2) .xls : 원본 유지 + report.xlsx 생성 + zip으로 묶기
    # -----------------------------
    if ext == ".xls":
        zip_name = f"validation_{uid}.zip"
        output_zip = os.path.join(output_dir, zip_name)

        # zip에 넣을 파일들
        base_xls_name = f"original_{uid}.xls"
        report_xlsx_name = f"report_{uid}.xlsx"

        temp_xls_path = os.path.join(output_dir, base_xls_name)
        temp_report_path = os.path.join(output_dir, report_xlsx_name)

        try:
            # 1) 원본 .xls를 그대로 복사(형식 유지)
            shutil.copy2(original_file, temp_xls_path)

            # 2) .xls에서 "원본 행의 값"을 뽑아오기 위해 raw로 읽기 (header=None)
            #    ※ .xls 읽으려면 xlrd가 필요해. (README에 pip install xlrd==2.0.1 권장)
            raw = pd.read_excel(original_file, sheet_name=0, header=None, engine="xlrd")

            # 헤더 행 탐지(0-based)
            header_row0 = 0
            max_scan = min(50, len(raw))
            for r in range(max_scan):
                row_vals = [str(v or "").strip() for v in raw.iloc[r].tolist()]
                if any("사원번호" in v for v in row_vals):
                    header_row0 = r
                    break

            headers = [str(v or "").strip().replace("\n", " ") for v in raw.iloc[header_row0].tolist()]
            df = raw.iloc[header_row0 + 1 :].copy()
            df.columns = headers

            # df의 실제 데이터 시작 엑셀 행번호(1-based)
            data_start_excel_row = (header_row0 + 1) + 1  # header_row0(0-based) + header(1줄) + 엑셀 1-based

            # col_map (컬럼명 -> df에서의 열 index)
            df_col_map = {str(c).strip(): c for c in df.columns}

            def find_df_col(keywords):
                # 정확 일치
                for kw in keywords:
                    kw = str(kw).strip()
                    if kw in df_col_map:
                        return kw
                # 포함 매칭
                for kw in keywords:
                    kw = str(kw).strip()
                    for c in df.columns:
                        if kw in str(c):
                            return c
                return None

            c_emp = find_df_col(["사원번호", "사번", "ID"])
            c_hire = find_df_col(["입사일자", "입사일", "채용일"])
            c_pay = find_df_col(["기준급여", "기본급"])
            c_birth = find_df_col(["생년월일", "출생", "생일"])

            # 3) report.xlsx 생성 (오류리포트 + AI_REPORT)
            wb = openpyxl.Workbook()
            # 기본 시트 제거
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]

            report_ws = wb.create_sheet("오류리포트", 0)
            headers_r = [
                "No", "원본행",
                "사원번호", "생년월일", "입사일자", "기준급여",
                "오류컬럼", "오류메시지"
            ]
            for j, h in enumerate(headers_r, start=1):
                cell = report_ws.cell(row=1, column=j, value=h)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")

            out_r = 2
            for i, err in enumerate(python_errors or [], start=1):
                row_idx = err.get("row")
                if not row_idx:
                    continue
                try:
                    excel_row = int(row_idx)
                except:
                    continue

                # excel_row -> df index로 변환
                df_index = excel_row - data_start_excel_row
                row_vals = None
                if 0 <= df_index < len(df):
                    row_vals = df.iloc[df_index]

                report_ws.cell(out_r, 1, i)
                report_ws.cell(out_r, 2, excel_row)

                report_ws.cell(out_r, 3, row_vals[c_emp] if (row_vals is not None and c_emp) else None)
                report_ws.cell(out_r, 4, row_vals[c_birth] if (row_vals is not None and c_birth) else None)
                report_ws.cell(out_r, 5, row_vals[c_hire] if (row_vals is not None and c_hire) else None)
                report_ws.cell(out_r, 6, row_vals[c_pay] if (row_vals is not None and c_pay) else None)

                report_ws.cell(out_r, 7, _norm(err.get("column")))
                report_ws.cell(out_r, 8, _err_msg(err))

                out_r += 1

            _autosize(report_ws)

            ai_ws = wb.create_sheet("AI_REPORT")
            ai_ws["A1"] = "AI 오류 분석 리포트"
            ai_ws["A1"].font = Font(bold=True, size=14)
            ai_ws["A3"] = ai_report or "(리포트 없음)"
            ai_ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
            ai_ws.column_dimensions["A"].width = 120
            ai_ws.row_dimensions[3].height = 600

            wb.save(temp_report_path)

            # 4) zip으로 묶어서 반환
            with zipfile.ZipFile(output_zip, "w", compression=zipfile.ZIP_DEFLATED) as zf:
                zf.write(temp_xls_path, arcname=base_xls_name)
                zf.write(temp_report_path, arcname=report_xlsx_name)

            # 임시파일 정리(원하면 지워도 되고 남겨도 됨)
            # os.remove(temp_xls_path)
            # os.remove(temp_report_path)

            return output_zip

        except Exception as e:
            print("❌ exporter(xls) error:", e)
            return None

    # 그 외 확장자
    print("❌ exporter error: unsupported file type:", ext)
    return None
# backend/exporter.py
import os
import uuid
import shutil
import openpyxl
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.utils import get_column_letter
import pandas as pd
import re


def _clean_text(x) -> str:
    s = "" if pd.isna(x) else str(x)
    s = s.replace("\n", " ").strip()
    s = re.sub(r"\s+", " ", s)
    return s


def _find_header_row_in_df(raw: pd.DataFrame) -> int:
    """
    위쪽 50줄에서 헤더 후보 행 찾기:
    '사원번호' + ('생년월일' or '입사일자/입사일') 중 하나가 있는 행
    """
    must = ["사원번호"]
    should = ["생년월일", "입사일", "입사일자"]

    max_scan = min(50, len(raw))
    for r in range(max_scan):
        row = raw.iloc[r].tolist()
        row_strs = [_clean_text(v) for v in row]

        if not any(m in cell for m in must for cell in row_strs):
            continue
        if not any(any(s in cell for cell in row_strs) for s in should):
            continue
        return r
    return 0


def _pick_sheet_from_xls(xls: pd.ExcelFile) -> str | int:
    # 1) 정확 일치
    for name in xls.sheet_names:
        if _clean_text(name) == "(2-2) 재직자 명부":
            return name
    # 2) 키워드 포함
    for name in xls.sheet_names:
        norm = _clean_text(name)
        if "재직" in norm and "명부" in norm:
            return name
    # 3) 내용으로 탐지: '사원번호'가 있는 시트 찾기
    for name in xls.sheet_names:
        try:
            raw = pd.read_excel(xls, sheet_name=name, header=None)
            # 상단 30줄만 빠르게 검사
            top = raw.head(30).astype(str).fillna("")
            if top.apply(lambda col: col.str.contains("사원번호").any()).any():
                return name
        except:
            pass
    return 0


def _autosize_worksheet(ws, max_width=45):
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for row in range(1, min(ws.max_row, 2000) + 1):
            v = ws.cell(row=row, column=col).value
            if v is None:
                continue
            v = str(v)
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[letter].width = min(max(10, max_len + 2), max_width)


def _make_report_sheet(wb, sheet_name: str, python_errors, df_for_values: pd.DataFrame | None):
    # 기존 있으면 삭제
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    # 헤더
    headers = ["엑셀행", "사원번호", "오류컬럼", "값", "오류메시지"]
    ws.append(headers)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")

    # 내용
    for err in (python_errors or []):
        r = err.get("row")
        col = err.get("column")
        msg = err.get("msg") or err.get("message") or ""
        emp_id = err.get("id") or ""

        value = ""
        if df_for_values is not None and isinstance(r, int) and r >= 2:
            df_idx = r - 2  # rules.py가 row_num = index+2라서
            if 0 <= df_idx < len(df_for_values) and col in df_for_values.columns:
                try:
                    value = df_for_values.iloc[df_idx][col]
                except:
                    value = ""
        ws.append([r, emp_id, col, value, msg])

    # 보기 좋게
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{ws.max_row}"
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    _autosize_worksheet(ws, max_width=60)
    return ws


def save_validation_results(original_file: str, python_errors, ai_report: str, output_dir: str = "temp_uploads"):
    """
    결과:
    - .xlsx 입력: 원본 서식 유지 + 오류 셀 하이라이트 + AI_REPORT/ERROR_REPORT 추가
    - .xls 입력: '명부 시트'를 자동 탐지해 깔끔한 .xlsx로 재생성 + 오류 셀 하이라이트 + 리포트 시트 추가
    """
    os.makedirs(output_dir, exist_ok=True)
    result_name = f"validation_{uuid.uuid4().hex[:10]}.xlsx"
    output_path = os.path.join(output_dir, result_name)

    try:
        error_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        # =========================
        # 1) 원본이 .xlsx인 경우: 복사 후 서식 유지
        # =========================
        if original_file.lower().endswith(".xlsx"):
            shutil.copy2(original_file, output_path)
            wb = openpyxl.load_workbook(output_path)

            # 대상 시트 선택
            if "(2-2) 재직자 명부" in wb.sheetnames:
                ws = wb["(2-2) 재직자 명부"]
            elif "(2-1) 명부" in wb.sheetnames:
                ws = wb["(2-1) 명부"]
            else:
                ws = wb.active

            # 헤더 행 찾기
            header_row = 1
            for r in range(1, min(50, ws.max_row) + 1):
                row_values = [str(ws.cell(row=r, column=c).value or "") for c in range(1, ws.max_column + 1)]
                if any("사원번호" in v for v in row_values):
                    header_row = r
                    break

            # 컬럼명 -> 컬럼 인덱스
            col_map = {}
            for c in range(1, ws.max_column + 1):
                key = ws.cell(row=header_row, column=c).value
                if key is not None:
                    col_map[str(key).replace("\n", " ").strip()] = c

            # 하이라이트 (rules.py가 row_num을 실제 엑셀 행으로 준다고 가정)
            for err in (python_errors or []):
                row_idx = err.get("row")
                col_name = err.get("column")
                if not row_idx or not col_name:
                    continue
                excel_row = int(row_idx)
                col_idx = col_map.get(str(col_name).strip())
                if col_idx:
                    ws.cell(row=excel_row, column=col_idx).fill = error_fill

            # AI_REPORT
            if "AI_REPORT" in wb.sheetnames:
                del wb["AI_REPORT"]
            rws = wb.create_sheet("AI_REPORT")
            rws["A1"] = "AI 오류 분석 리포트"
            rws["A1"].font = Font(bold=True)
            rws["A2"] = ai_report or "(리포트 없음)"
            rws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
            rws.column_dimensions["A"].width = 120

            # ERROR_REPORT (값을 뽑기 위해서는 df가 필요하지만, 여기서는 비워도 OK)
            _make_report_sheet(wb, "ERROR_REPORT", python_errors, df_for_values=None)

            wb.save(output_path)
            return output_path

        # =========================
        # 2) 원본이 .xls인 경우: "명부 시트" 자동 탐지 후, 보기 좋게 새로 생성
        # =========================
        else:
            # xls 로딩: 시트 선택 + 헤더 탐지 + 데이터프레임 생성
            xls = pd.ExcelFile(original_file, engine="xlrd")
            target_sheet = _pick_sheet_from_xls(xls)
            raw = pd.read_excel(xls, sheet_name=target_sheet, header=None)

            header_row = _find_header_row_in_df(raw)
            header = [_clean_text(v) for v in raw.iloc[header_row].tolist()]

            df = raw.iloc[header_row + 1 :].copy()
            df.columns = header
            df = df.dropna(axis=0, how="all").dropna(axis=1, how="all")
            df.columns = [str(c).replace("\n", " ").strip() for c in df.columns]

            # 새 워크북 생성
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "(2-2) 재직자 명부"  # 일단 통일

            # 헤더 작성
            ws.append(list(df.columns))
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=c)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
                cell.alignment = Alignment(wrap_text=True, vertical="center")

            # 데이터 작성
            for _, row in df.iterrows():
                ws.append(list(row.values))

            # 보기 좋게: 고정/필터/줄바꿈/열너비
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
            for r in range(2, ws.max_row + 1):
                for c in range(1, ws.max_column + 1):
                    ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")

            _autosize_worksheet(ws, max_width=35)

            # 오류 셀 하이라이트
            col_map = {str(name).strip(): idx + 1 for idx, name in enumerate(df.columns)}
            for err in (python_errors or []):
                row_idx = err.get("row")   # rules.py가 row_num = index+2
                col_name = err.get("column")
                if not row_idx or not col_name:
                    continue
                excel_row = int(row_idx)   # 생성한 시트는 헤더가 1행이라 그대로 맞음
                col_idx = col_map.get(str(col_name).strip())
                if col_idx and 2 <= excel_row <= ws.max_row:
                    ws.cell(row=excel_row, column=col_idx).fill = error_fill

            # AI_REPORT
            rws = wb.create_sheet("AI_REPORT")
            rws["A1"] = "AI 오류 분석 리포트"
            rws["A1"].font = Font(bold=True)
            rws["A2"] = ai_report or "(리포트 없음)"
            rws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
            rws.column_dimensions["A"].width = 120

            # ERROR_REPORT (값 포함)
            _make_report_sheet(wb, "ERROR_REPORT", python_errors, df_for_values=df)

            wb.save(output_path)
            return output_path

    except Exception as e:
        print("❌ exporter error:", e)
        return None
